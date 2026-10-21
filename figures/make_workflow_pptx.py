#!/usr/bin/env python3
"""Draw the DANTE_TIR principle scheme as an editable PowerPoint slide.

Layout and colours follow the DANTE_LTR workflow figure. Coordinates are
given in pixels of a 1800 x 1050 canvas (1 px = 8128 EMU, slide 16 in wide).

Usage:
    python make_workflow_pptx.py [output.pptx]
"""
import random
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

W, H = 1800, 1050
EMU_PER_PX = 8128
PT_PER_PX = 0.64
FONT = "DejaVu Sans"

# colours (Okabe-Ito based, as in the DANTE_LTR figure)
HAT = "0072B2"
MUDR = "CC79A7"
CACTA = "009E73"
CACTA_TINT = "CCEBE0"
TIR = "333333"
TSD = "E69F00"
TRACK = "CCCCCC"
WINDOW_FILL = "EBEBEB"
WINDOW_LINE = "AAAAAA"
BOUNDARY = "C0392B"
OK = "1B7F3B"
TEXT = "111111"
MUTED = "6E6E6E"
NUMBER = "777777"

# x coordinates shared by rows 2-4 (one CACTA locus)
DOM_X0, DOM_X1 = 900, 1030
WIN_L0, WIN_R1 = 380, 1550
END_5, END_3 = 560, 1370  # element boundaries
TSD_W, TIR_W = 28, 100


def px(v):
    return Emu(int(round(v * EMU_PER_PX)))


def rgb(hex_):
    return RGBColor.from_string(hex_)


class Slide:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = px(W)
        self.prs.slide_height = px(H)
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bg = self.slide.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb("FFFFFF")
        self.shapes = self.slide.shapes

    # --- primitives -------------------------------------------------------
    def _style(self, shp, fill=None, line=None, line_w=1.5, dash=None):
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(line_w)
            if dash:
                shp.line.dash_style = dash
        _drop_style(shp)
        return shp

    def rect(self, x, y, w, h, fill=None, line=None, line_w=1.5, dash=None,
             shape=MSO_SHAPE.RECTANGLE):
        shp = self.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
        return self._style(shp, fill, line, line_w, dash)

    def text(self, x, y, w, h, runs, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.MIDDLE, shp=None):
        """runs: list of (text, size_px, colour, bold)."""
        if shp is None:
            shp = self.shapes.add_textbox(px(x), px(y), px(w), px(h))
        tf = shp.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        p.alignment = align
        for txt, size, colour, bold in runs:
            r = p.add_run()
            r.text = txt
            r.font.name = FONT
            r.font.size = Pt(size * PT_PER_PX)
            r.font.color.rgb = rgb(colour)
            r.font.bold = bold
        return shp

    def line(self, x0, y0, x1, y1, colour=TRACK, w=2.0, dash=None,
             head=False, tail=False):
        c = self.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      px(x0), px(y0), px(x1), px(y1))
        c.line.color.rgb = rgb(colour)
        c.line.width = Pt(w)
        if dash:
            c.line.dash_style = dash
        _arrowheads(c.line._get_or_add_ln(), head, tail)
        _drop_style(c)
        return c

    def path(self, start, segments, colour, w=2.0, tail=True):
        """Open path from start; segments are ("L", p) or ("C", c1, c2, p)."""
        pts = [start] + [q for seg in segments for q in seg[1:]]
        x0 = min(q[0] for q in pts)
        y0 = min(q[1] for q in pts)
        x1 = max(q[0] for q in pts)
        y1 = max(q[1] for q in pts)
        shp = self.rect(x0, y0, x1 - x0, y1 - y0)
        spPr = shp._element.spPr
        spPr.remove(spPr.find(qn("a:prstGeom")))

        def pt(parent, q):
            el = etree.SubElement(parent, qn("a:pt"))
            el.set("x", str(px(q[0] - x0)))
            el.set("y", str(px(q[1] - y0)))

        geom = etree.Element(qn("a:custGeom"))
        for tag in ("a:avLst", "a:gdLst", "a:ahLst", "a:cxnLst"):
            etree.SubElement(geom, qn(tag))
        rect = etree.SubElement(geom, qn("a:rect"))
        for k, v in (("l", "0"), ("t", "0"), ("r", "r"), ("b", "b")):
            rect.set(k, v)
        lst = etree.SubElement(geom, qn("a:pathLst"))
        path = etree.SubElement(lst, qn("a:path"))
        path.set("w", str(px(x1 - x0)))
        path.set("h", str(px(y1 - y0)))
        path.set("fill", "none")
        pt(etree.SubElement(path, qn("a:moveTo")), start)
        for seg in segments:
            tag = "a:lnTo" if seg[0] == "L" else "a:cubicBezTo"
            el = etree.SubElement(path, qn(tag))
            for q in seg[1:]:
                pt(el, q)
        spPr.find(qn("a:xfrm")).addnext(geom)
        shp.line.color.rgb = rgb(colour)
        shp.line.width = Pt(w)
        _arrowheads(shp.line._get_or_add_ln(), False, tail)
        return shp

    def arrow(self, x, y, w, h, fill, label=None, left=False, size=19):
        """Pentagon arrow; left=True mirrors it, label stays upright."""
        shp = self.rect(x, y, w, h, fill=fill, shape=MSO_SHAPE.PENTAGON)
        if left:
            shp._element.spPr.find(qn("a:xfrm")).set("flipH", "1")
        if label:
            tip = h * 0.5
            lx = x + tip if left else x
            self.text(lx, y, w - tip, h, [(label, size, "FFFFFF", True)],
                      align=PP_ALIGN.CENTER)
        return shp

    # --- composite pieces -------------------------------------------------
    def row_title(self, y, num, title, caption, status=None):
        self.text(40, y, 40, 30, [(str(num), 24, NUMBER, False)])
        self.text(80, y, 1300, 30, [(title, 24, TEXT, False),
                                    ("      " + caption, 19, MUTED, False)])
        if status:
            self.text(1300, y, 460, 30, [(status, 20, OK, True)],
                      align=PP_ALIGN.RIGHT)

    def track(self, y, x0=40, x1=1760):
        self.line(x0, y, x1, y, TRACK, 1.5)

    def window(self, x0, x1, y, h=70):
        self.rect(x0, y - h / 2, x1 - x0, h, fill=WINDOW_FILL,
                  line=WINDOW_LINE, line_w=1.2,
                  dash=MSO_LINE_DASH_STYLE.DASH)

    def dimension(self, x0, x1, y, label):
        self.line(x0, y, x1, y, MUTED, 1.0, head=True, tail=True)
        self.text(x0, y - 26, x1 - x0, 22, [(label, 17, MUTED, False)],
                  align=PP_ALIGN.CENTER)

    def element(self, y, x5, x3, h=34, seq=True, faded=False):
        """TSD, 5'TIR, element body with domain, 3'TIR, TSD."""
        body_h = 14
        self.rect(x5 + TIR_W, y - body_h / 2, x3 - x5 - 2 * TIR_W, body_h,
                  fill=CACTA_TINT)
        self.arrow(DOM_X0, y - h / 2, DOM_X1 - DOM_X0, h, CACTA, "CACTA")
        self.rect(x5 - TSD_W, y - h / 2, TSD_W, h, fill=TSD)
        self.rect(x3, y - h / 2, TSD_W, h, fill=TSD)
        self.arrow(x5, y - h / 2, TIR_W, h, TIR, "5'TIR", size=17)
        self.arrow(x3 - TIR_W, y - h / 2, TIR_W, h, TIR, "3'TIR", left=True,
                   size=17)


def _drop_style(shp):
    style = shp._element.find(qn("p:style"))
    if style is not None:
        shp._element.remove(style)


def _arrowheads(ln, head, tail):
    for flag, tag in ((head, "a:headEnd"), (tail, "a:tailEnd")):
        if flag:
            el = etree.SubElement(ln, qn(tag))
            el.set("type", "triangle")
            el.set("w", "med")
            el.set("len", "med")


def mosaic(s, x0, x1, y, h, rnd):
    """Grey blocks of random shade: unrelated sequence at insertion sites."""
    x = x0
    while x < x1:
        w = min(rnd.randint(8, 26), x1 - x)
        g = rnd.randint(0xC4, 0xEE)
        s.rect(x, y, w, h, fill=f"{g:02X}{g:02X}{g:02X}")
        x += w


def main(out):
    rnd = random.Random(7)
    s = Slide()

    # 1 -- DANTE domains, several superfamilies
    s.row_title(45, 1, "DANTE transposase domains",
                "colour = the superfamily DANTE assigned")
    y = 120
    s.track(y)
    for x, col, lab in ((250, HAT, "hAT"), (520, CACTA, "CACTA"),
                        (900, MUDR, "MuDR"), (1180, HAT, "hAT"),
                        (1460, CACTA, "CACTA")):
        s.arrow(x, y - 17, 130, 34, col, lab)

    # 2 -- flanking windows of one domain
    s.row_title(195, 2, "Flanking regions",
                "6 kb on each side of every domain, "
                "processed separately for each superfamily")
    y = 285
    s.track(y)
    s.window(WIN_L0, DOM_X0 - 10, y)
    s.window(DOM_X1 + 10, WIN_R1, y)
    s.arrow(DOM_X0, y - 17, DOM_X1 - DOM_X0, 34, CACTA, "CACTA")
    s.dimension(WIN_L0 + 10, DOM_X0 - 20, y + 12, "6 kb")
    s.dimension(DOM_X1 + 20, WIN_R1 - 10, y + 12, "6 kb")

    # 3 -- copies of one superfamily compared
    s.row_title(370, 3, "Copies compared",
                "flanks of all copies are aligned; the element is shared, "
                "the insertion site is not",
                status="✓ element ends")
    top, step, h = 425, 22, 14
    n = 5
    for i in range(n):
        yy = top + i * step
        mosaic(s, WIN_L0, END_5, yy, h, rnd)
        s.rect(END_5, yy, END_3 - END_5, h, fill=CACTA_TINT)
        s.arrow(DOM_X0, yy, DOM_X1 - DOM_X0, h, CACTA)
        mosaic(s, END_3, WIN_R1, yy, h, rnd)
    bottom = top + (n - 1) * step + h
    for x in (END_5, END_3):
        s.line(x, top - 8, x, bottom + 8, BOUNDARY, 2.5)
    for x0, x1 in ((WIN_L0, END_5), (END_3, WIN_R1)):
        s.text(x0 - 60, bottom + 12, x1 - x0 + 120, 22,
               [("insertion sites (masked)", 16, MUTED, False)],
               align=PP_ALIGN.CENTER)
    s.text(END_5, bottom + 12, END_3 - END_5, 22,
           [("similarity ends at the element boundary", 16, BOUNDARY,
             False)], align=PP_ALIGN.CENTER)

    # 4 -- TIRs and TSDs at the boundaries
    s.row_title(590, 4, "TIR and TSD",
                "element ends form an inverted repeat, "
                "flanked by a target site duplication",
                status="✓ complete element")
    y = 665
    s.track(y, x1=1440)
    s.element(y, END_5, END_3)
    for x in (END_5 - TSD_W, END_3):
        s.text(x - 30, y + 24, TSD_W + 60, 22, [("TSD", 17, TSD, True)],
               align=PP_ALIGN.CENTER)
    s.text(END_5, y - 46, 200, 22,
           [("CACTACG…", 17, TEXT, False)])
    s.text(END_3 - 200, y - 46, 200, 22,
           [("…CGTAGTG", 17, TEXT, False)], align=PP_ALIGN.RIGHT)
    # key: TSD expected by the superfamily-specific detector
    kx0, kx1, ky = 1490, 1760, 628
    s.text(kx0, ky, kx1 - kx0, 18,
           [("TSD expected per superfamily", 15, MUTED, False)])
    s.line(kx0, ky + 22, kx1, ky + 22, TRACK, 1.0)
    for i, (name, col, tsd) in enumerate((
            ("Tc1/Mariner", TEXT, "2 bp (TA)"),
            ("PIF/Harbinger", TEXT, "3 bp"),
            ("CACTA", CACTA, "3 bp"),
            ("hAT", HAT, "\u2265 6 bp"),
            ("MuDR", MUDR, "\u2265 7 bp"))):
        yy = ky + 28 + i * 19
        s.text(kx0, yy, 160, 17, [(name, 15, col, True)])
        s.text(kx0 + 150, yy, kx1 - kx0 - 150, 17,
               [(tsd, 15, TEXT, False)], align=PP_ALIGN.RIGHT)

    # 5 -- iterative search with the ends found so far
    s.row_title(775, 5, "Iterative search",
                "detected ends are searched for in the flanks of copies "
                "not yet resolved",
                status="↻ repeated rounds")
    lib_x, lib_y, lib_w, lib_h = 70, 845, 250, 110
    s.rect(lib_x, lib_y, lib_w, lib_h, fill="F4F4F4", line=WINDOW_LINE,
           line_w=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    s.text(lib_x, lib_y + 8, lib_w, 22,
           [("detected element ends", 16, MUTED, False)],
           align=PP_ALIGN.CENTER)
    for i in range(3):
        yy = lib_y + 42 + i * 22
        s.arrow(lib_x + 45, yy, 50, 14, TIR)
        s.arrow(lib_x + 155, yy, 50, 14, TIR, left=True)
        s.line(lib_x + 95, yy + 7, lib_x + 155, yy + 7, TRACK, 1.5)
    s.line(lib_x + lib_w + 10, lib_y + lib_h / 2, 440, lib_y + lib_h / 2,
           MUTED, 2.0, tail=True)

    # two loci where the search adds new elements
    loci = ((860, 610, 1300), (935, 520, 1400))
    for yy, x5, x3 in loci:
        s.line(460, yy, 1610, yy, TRACK, 1.5)
        s.window(470, DOM_X0 - 10, yy, h=54)
        s.window(DOM_X1 + 10, 1600, yy, h=54)
        s.element(yy, x5, x3, h=28)
    # return loop back to the library
    loop_y = 995
    lib_cx = lib_x + lib_w / 2
    mid = (loci[0][0] + loci[1][0]) / 2
    s.path((1612, mid),
           [("C", (1700, mid), (1700, loop_y), (1610, loop_y)),
            ("L", (lib_cx + 50, loop_y)),
            ("C", (lib_cx, loop_y), (lib_cx, loop_y),
             (lib_cx, lib_y + lib_h + 6))],
           MUTED, 2.0)
    s.text(460, loop_y + 6, 1150, 24,
           [("new elements are added to the set of ends and the search "
             "is repeated", 16, MUTED, False)], align=PP_ALIGN.CENTER)

    s.prs.save(out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "dante_tir_workflow.pptx")
